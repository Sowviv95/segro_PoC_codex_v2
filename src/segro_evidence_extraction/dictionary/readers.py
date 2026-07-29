"""Dictionary file readers for XLSX and CSV inputs."""

from __future__ import annotations

import csv
from collections import defaultdict
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from segro_evidence_extraction.dictionary.models import (
    DictionaryFormat,
    DictionaryMetadata,
    RawDictionary,
    RawDictionaryRow,
    source_name,
)


class DictionaryReadError(ValueError):
    """Raised when a dictionary file cannot be read as the requested format."""


class MissingWorksheetError(DictionaryReadError):
    """Raised when an XLSX worksheet is missing."""


class DictionaryReader:
    """Generic dictionary reader interface."""

    def read(self, path: Path, *, sheet_name: str | None, header_row: int) -> RawDictionary:
        raise NotImplementedError

    def inspect(self, path: Path, *, sheet_name: str | None, header_row: int) -> RawDictionary:
        return self.read(path, sheet_name=sheet_name, header_row=header_row)


def read_dictionary(path: Path, *, sheet_name: str | None, header_row: int = 1) -> RawDictionary:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return XlsxDictionaryReader().read(path, sheet_name=sheet_name, header_row=header_row)
    if suffix == ".csv":
        return CsvDictionaryReader().read(path, sheet_name=sheet_name, header_row=header_row)
    msg = f"Unsupported dictionary format: {path.suffix}"
    raise DictionaryReadError(msg)


def available_sheets(path: Path) -> list[str]:
    if path.suffix.lower() != ".xlsx":
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


class XlsxDictionaryReader(DictionaryReader):
    """Read XLSX dictionaries with physical row preservation."""

    def read(self, path: Path, *, sheet_name: str | None, header_row: int) -> RawDictionary:
        start = perf_counter()
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (OSError, InvalidFileException) as exc:
            msg = f"Malformed XLSX dictionary: {path}"
            raise DictionaryReadError(msg) from exc
        file_read_ms = (perf_counter() - start) * 1000
        try:
            selected_sheet = sheet_name or workbook.sheetnames[0]
            if selected_sheet not in workbook.sheetnames:
                msg = f"Worksheet not found: {selected_sheet}"
                raise MissingWorksheetError(msg)
            worksheet = workbook[selected_sheet]
            sheet_names = list(workbook.sheetnames)
            rows_iter = worksheet.iter_rows(values_only=True)
            rows = list(rows_iter)
        finally:
            workbook.close()
        return _raw_dictionary_from_rows(
            path,
            rows,
            dictionary_format=DictionaryFormat.XLSX,
            source=selected_sheet,
            sheet_names=sheet_names,
            header_row=header_row,
            file_read_ms=file_read_ms,
        )


class CsvDictionaryReader(DictionaryReader):
    """Read CSV dictionaries with physical row preservation."""

    def read(self, path: Path, *, sheet_name: str | None, header_row: int) -> RawDictionary:
        _ = sheet_name
        start = perf_counter()
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = [tuple(row) for row in csv.reader(handle)]
        except OSError as exc:
            msg = f"Malformed CSV dictionary: {path}"
            raise DictionaryReadError(msg) from exc
        file_read_ms = (perf_counter() - start) * 1000
        return _raw_dictionary_from_rows(
            path,
            rows,
            dictionary_format=DictionaryFormat.CSV,
            source=source_name(path, None),
            sheet_names=[],
            header_row=header_row,
            file_read_ms=file_read_ms,
        )


def _raw_dictionary_from_rows(
    path: Path,
    rows: Sequence[Sequence[object]],
    *,
    dictionary_format: DictionaryFormat,
    source: str,
    sheet_names: list[str],
    header_row: int,
    file_read_ms: float,
) -> RawDictionary:
    _ = file_read_ms
    if header_row < 1 or header_row > len(rows):
        msg = f"Header row {header_row} is outside the file row range"
        raise DictionaryReadError(msg)
    header_values = [_cell_to_header(value) for value in rows[header_row - 1]]
    columns = _make_unique_headers(header_values)
    raw_rows: list[RawDictionaryRow] = []
    for physical_row_number, row_values in enumerate(rows[header_row:], start=header_row + 1):
        padded = list(row_values) + [None] * max(len(columns) - len(row_values), 0)
        values = {column: _clean_cell(padded[index]) for index, column in enumerate(columns)}
        raw_rows.append(RawDictionaryRow(physical_row_number=physical_row_number, values=values))
    duplicate_or_blank = _duplicate_or_blank_columns(header_values)
    mixed_types = _mixed_type_columns(columns, rows[header_row:])
    metadata = DictionaryMetadata(
        dictionary_path=str(path),
        dictionary_filename=path.name,
        dictionary_id=path.stem,
        source_name=source,
        format=dictionary_format,
        sheet_names=sheet_names,
        selected_sheet=source if dictionary_format == DictionaryFormat.XLSX else None,
        header_row=header_row,
        columns=columns,
        mapping_version="unmapped",
        ingested_at=datetime.now(UTC),
    )
    return RawDictionary(
        metadata=metadata,
        rows=raw_rows,
        duplicate_or_blank_columns=duplicate_or_blank,
        mixed_type_columns=mixed_types,
    )


def _cell_to_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _make_unique_headers(headers: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers, start=1):
        base = header or f"blank_column_{index}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return result


def _duplicate_or_blank_columns(headers: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    for header in headers:
        counts[header] = counts.get(header, 0) + 1
    issues = [header for header, count in counts.items() if header == "" or count > 1]
    return sorted(set(issues))


def _mixed_type_columns(
    columns: list[str],
    rows: Sequence[Sequence[object]],
) -> dict[str, list[str]]:
    types_by_column: dict[str, set[str]] = defaultdict(set)
    for row_values in rows:
        for index, column in enumerate(columns):
            if index >= len(row_values):
                continue
            value = row_values[index]
            if value is not None and str(value).strip() != "":
                types_by_column[column].add(type(value).__name__)
    return {
        column: sorted(types)
        for column, types in types_by_column.items()
        if len(types) > 1
    }


def _clean_cell(value: object) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, str):
        normalized = " ".join(value.replace("\r", "\n").split())
        return normalized or None
    if isinstance(value, bool | int | float):
        return value
    return str(value)
