"""Cheap source metadata inspection."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from segro_evidence_extraction.models.source import FileType

MetadataValue = str | int | float | bool | None


def inspect_metadata(path: Path, file_type: FileType) -> tuple[dict[str, MetadataValue], list[str]]:
    if file_type == FileType.PDF:
        return _pdf_metadata(path)
    if file_type == FileType.XLSX:
        return _xlsx_metadata(path)
    if file_type == FileType.CSV:
        return _csv_metadata(path)
    if file_type == FileType.TEXT:
        return _text_metadata(path)
    if file_type == FileType.IMAGE:
        return _image_metadata(path)
    return {}, []


def _pdf_metadata(path: Path) -> tuple[dict[str, str | int | float | bool | None], list[str]]:
    warnings: list[str] = []
    try:
        page_count = _stream_pdf_page_count(path)
        first_header = _read_bounded_header(path, 4096)
    except OSError as exc:
        return {}, [f"PDF metadata read failed: {exc}"]
    if page_count == 0:
        warnings.append("PDF page count could not be inferred from lightweight metadata.")
    return {"page_count": page_count or None, "header_text": first_header[:500]}, warnings


def _stream_pdf_page_count(path: Path) -> int:
    count = 0
    carry = b""
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            data = carry + block
            count += data.count(b"/Type /Page")
            count += data.count(b"/Type/Page")
            carry = data[-32:]
    return count


def _read_bounded_header(path: Path, size: int) -> str:
    with path.open("rb") as handle:
        return handle.read(size).decode("latin-1", errors="ignore")


def _xlsx_metadata(path: Path) -> tuple[dict[str, str | int | float | bool | None], list[str]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - metadata inspection should not abort ingestion.
        return {}, [f"Workbook metadata read failed: {exc}"]
    try:
        sheet_names = list(workbook.sheetnames)
    finally:
        workbook.close()
    return {"sheet_count": len(sheet_names), "sheet_names": "|".join(sheet_names)}, []


def _csv_metadata(path: Path) -> tuple[dict[str, str | int | float | bool | None], list[str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader, [])
    except OSError as exc:
        return {}, [f"CSV metadata read failed: {exc}"]
    return {"column_count": len(header), "columns": "|".join(header[:50])}, []


def _text_metadata(path: Path) -> tuple[dict[str, str | int | float | bool | None], list[str]]:
    try:
        sample = path.read_text(encoding="utf-8", errors="ignore")[:500]
    except OSError as exc:
        return {}, [f"Text metadata read failed: {exc}"]
    return {"header_text": sample}, []


def _image_metadata(path: Path) -> tuple[dict[str, str | int | float | bool | None], list[str]]:
    return {"image_filename": path.name}, []
