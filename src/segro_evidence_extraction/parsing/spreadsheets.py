"""Incremental spreadsheet and CSV parsing."""

from __future__ import annotations

import csv
import time
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from segro_evidence_extraction.models.classification import ClassificationSubjectType
from segro_evidence_extraction.models.common import ProvenanceRef
from segro_evidence_extraction.models.source import SourceRegistryEntry
from segro_evidence_extraction.parsing.cache import JsonParseCache, cache_key, config_hash
from segro_evidence_extraction.parsing.classification import (
    CLASSIFIER_VERSION,
    classify_page_evidence,
)
from segro_evidence_extraction.parsing.interfaces import ProgressCallback
from segro_evidence_extraction.parsing.models import (
    NORMALIZATION_VERSION,
    DocumentTiming,
    ParsedSheet,
    ParseOptions,
    ParserWarning,
    ParsingConfig,
    ParsingStatus,
)
from segro_evidence_extraction.parsing.normalization import normalize_text, truncate_text
from segro_evidence_extraction.parsing.quality import assess_text_quality

XLSX_PARSER_NAME = "openpyxl-read-only"
XLSX_PARSER_VERSION = "openpyxl-read-only-v1"
CSV_PARSER_NAME = "csv-stream"
CSV_PARSER_VERSION = "csv-stream-v1"


class XlsxSheetParser:
    parser_name = XLSX_PARSER_NAME
    parser_version = XLSX_PARSER_VERSION

    def parse(
        self,
        source: SourceRegistryEntry,
        *,
        source_path: Path,
        config: ParsingConfig,
        options: ParseOptions,
        cache: JsonParseCache | None,
        progress: ProgressCallback | None,
    ) -> Iterator[ParsedSheet | ParserWarning | DocumentTiming]:
        doc_start = time.perf_counter()
        timing = DocumentTiming(
            source_id=source.source_id,
            logical_path=source.logical_path,
            parser_name=self.parser_name,
        )
        parsed = 0
        started = datetime.now(UTC)
        try:
            open_start = time.perf_counter()
            workbook = load_workbook(
                source_path,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            timing.open_ms = (time.perf_counter() - open_start) * 1000
            try:
                for sheet_number, worksheet in enumerate(workbook.worksheets, start=1):
                    if options.max_pages is not None and parsed >= options.max_pages:
                        break
                    label = f"sheet-{sheet_number}"
                    if progress:
                        progress("parse:sheet", f"{source.logical_path} Sheet {sheet_number}")
                    key = _sheet_cache_key(
                        source,
                        label,
                        config,
                        self.parser_name,
                        self.parser_version,
                    )
                    cached = cache.read_sheet(key) if cache and options.use_cache else None
                    if cached is not None:
                        cached.status = ParsingStatus.CACHE_HIT
                        parsed += 1
                        yield cached
                        continue
                    sheet_start = time.perf_counter()
                    preview_rows: list[str] = []
                    max_columns = 0
                    row_count = 0
                    iter_start = time.perf_counter()
                    for row in worksheet.iter_rows(values_only=False):
                        row_count += 1
                        values = [_cell_to_text(cell.value) for cell in row]
                        while values and values[-1] == "":
                            values.pop()
                        max_columns = max(max_columns, len(values))
                        if row_count <= config.spreadsheet_preview_rows:
                            preview_rows.append(",".join(values))
                    timing.iteration_ms += (time.perf_counter() - iter_start) * 1000
                    norm_start = time.perf_counter()
                    text, _truncated = truncate_text(
                        normalize_text("\n".join(preview_rows)),
                        config.max_extracted_chars_per_page,
                    )
                    timing.normalization_ms += (time.perf_counter() - norm_start) * 1000
                    quality_start = time.perf_counter()
                    quality = assess_text_quality(text)
                    timing.quality_ms += (time.perf_counter() - quality_start) * 1000
                    table_indicator = row_count > 0 and max_columns > 1
                    class_start = time.perf_counter()
                    classification = classify_page_evidence(
                        source_id=source.source_id,
                        page_or_sheet=label,
                        subject_type=ClassificationSubjectType.SHEET,
                        text=text,
                        quality=quality,
                        file_type="xlsx",
                        table_indicator=table_indicator,
                        drawing_indicator=False,
                    )
                    timing.classification_ms += (time.perf_counter() - class_start) * 1000
                    sheet = ParsedSheet(
                        sheet_id=f"{source.source_id}:s{sheet_number}",
                        source_id=source.source_id,
                        content_hash=source.file_hash,
                        sheet_number=sheet_number,
                        sheet_name=str(worksheet.title),
                        parser_name=self.parser_name,
                        parser_version=self.parser_version,
                        row_count=row_count,
                        column_count=max_columns,
                        preview_text=text,
                        character_count=quality.character_count,
                        word_count=quality.word_count,
                        table_indicator=table_indicator,
                        status=ParsingStatus.PARSED,
                        parsing_duration_ms=(time.perf_counter() - sheet_start) * 1000,
                        provenance=[
                            ProvenanceRef(
                                source_id=source.source_id,
                                page_or_sheet=label,
                                notes="Parsed from XLSX worksheet in read-only mode.",
                            )
                        ],
                        classification=classification,
                    )
                    if cache and options.use_cache:
                        cache_start = time.perf_counter()
                        cache.write_sheet(key, sheet)
                        timing.cache_write_ms += (time.perf_counter() - cache_start) * 1000
                    parsed += 1
                    yield sheet
            finally:
                workbook.close()
        except Exception as exc:  # noqa: BLE001
            timing.failures += 1
            yield ParserWarning(
                code="xlsx_parse_failed",
                message=f"XLSX parsing failed: {exc}",
                source_id=source.source_id,
            )
        finally:
            timing.pages_or_sheets = parsed
            timing.total_ms = (time.perf_counter() - doc_start) * 1000
            _ = started
            yield timing


class CsvSheetParser:
    parser_name = CSV_PARSER_NAME
    parser_version = CSV_PARSER_VERSION

    def parse(
        self,
        source: SourceRegistryEntry,
        *,
        source_path: Path,
        config: ParsingConfig,
        options: ParseOptions,
        cache: JsonParseCache | None,
        progress: ProgressCallback | None,
    ) -> Iterator[ParsedSheet | ParserWarning | DocumentTiming]:
        doc_start = time.perf_counter()
        timing = DocumentTiming(
            source_id=source.source_id,
            logical_path=source.logical_path,
            parser_name=self.parser_name,
        )
        label = "sheet-1"
        try:
            if progress:
                progress("parse:csv", f"{source.logical_path} streaming rows")
            key = _sheet_cache_key(source, label, config, self.parser_name, self.parser_version)
            cached = cache.read_sheet(key) if cache and options.use_cache else None
            if cached is not None:
                cached.status = ParsingStatus.CACHE_HIT
                yield cached
                return
            row_count = 0
            max_columns = 0
            preview_rows: list[str] = []
            open_start = time.perf_counter()
            with source_path.open(
                "r",
                encoding="utf-8-sig",
                errors="replace",
                newline="",
            ) as handle:
                timing.open_ms = (time.perf_counter() - open_start) * 1000
                iter_start = time.perf_counter()
                reader = csv.reader(handle)
                for row in reader:
                    row_count += 1
                    max_columns = max(max_columns, len(row))
                    if row_count <= config.csv_preview_rows:
                        preview_rows.append(",".join(row))
                timing.iteration_ms = (time.perf_counter() - iter_start) * 1000
            norm_start = time.perf_counter()
            text, _truncated = truncate_text(
                normalize_text("\n".join(preview_rows)),
                config.max_extracted_chars_per_page,
            )
            timing.normalization_ms = (time.perf_counter() - norm_start) * 1000
            quality_start = time.perf_counter()
            quality = assess_text_quality(text)
            timing.quality_ms = (time.perf_counter() - quality_start) * 1000
            class_start = time.perf_counter()
            classification = classify_page_evidence(
                source_id=source.source_id,
                page_or_sheet=label,
                subject_type=ClassificationSubjectType.SHEET,
                text=text,
                quality=quality,
                file_type="csv",
                table_indicator=max_columns > 1,
                drawing_indicator=False,
            )
            timing.classification_ms = (time.perf_counter() - class_start) * 1000
            sheet = ParsedSheet(
                sheet_id=f"{source.source_id}:s1",
                source_id=source.source_id,
                content_hash=source.file_hash,
                sheet_number=1,
                sheet_name=source.logical_path,
                parser_name=self.parser_name,
                parser_version=self.parser_version,
                row_count=row_count,
                column_count=max_columns,
                preview_text=text,
                character_count=quality.character_count,
                word_count=quality.word_count,
                table_indicator=max_columns > 1,
                status=ParsingStatus.PARSED,
                parsing_duration_ms=timing.total_ms,
                provenance=[
                    ProvenanceRef(
                        source_id=source.source_id,
                        page_or_sheet=label,
                        notes="Parsed from CSV using streaming rows.",
                    )
                ],
                classification=classification,
            )
            if cache and options.use_cache:
                cache.write_sheet(key, sheet)
            yield sheet
        except Exception as exc:  # noqa: BLE001
            timing.failures = 1
            yield ParserWarning(
                code="csv_parse_failed",
                message=f"CSV parsing failed: {exc}",
                source_id=source.source_id,
            )
        finally:
            timing.pages_or_sheets = 1 if timing.failures == 0 else 0
            timing.total_ms = (time.perf_counter() - doc_start) * 1000
            yield timing


def _sheet_cache_key(
    source: SourceRegistryEntry,
    label: str,
    config: ParsingConfig,
    parser_name: str,
    parser_version: str,
) -> str:
    return cache_key(
        source_id=source.source_id,
        content_hash=source.file_hash,
        page_or_sheet=label,
        parser_name=parser_name,
        parser_version=parser_version,
        normalization_version=NORMALIZATION_VERSION,
        classification_version=CLASSIFIER_VERSION,
        config_digest=config_hash(config),
    )


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
